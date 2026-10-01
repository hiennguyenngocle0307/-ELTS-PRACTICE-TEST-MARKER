# -*- coding: utf-8 -*-
"""Tạo đề Reading giữa kỳ RW B2.2 (docx + html cho pdf) từ 6 test nguồn (Drive: IELTS Reading Actual Tests).
python3 build_reading.py <thư mục output>   (HTML_OUT=<dir> để xuất HTML cho PDF)"""
import os,sys,re,html as _h,importlib
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE); sys.path.insert(0,os.path.join(HERE,'tests'))
_argv=sys.argv; sys.argv=[_argv[0]]
sys.path.insert(0,os.path.join(HERE,'..','_build')); import build as G
sys.argv=_argv
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
R,P,cellp,shade,shade_par,widths,noborder,rule,footer,part_heading,base_doc=G.R,G.P,G.cellp,G.shade,G.shade_par,G.widths,G.noborder,G.rule,G.footer,G.part_heading,G.base_doc
NAVY,PINK,LBLUE,RED,GREY=G.NAVY,G.PINK,G.LBLUE,G.RED,G.GREY
OUT=sys.argv[1] if len(sys.argv)>1 else os.path.join(HERE,'..','RW B2.2-MIDTERM TEST - READING')
YEAR=G.YEAR; COURSE="ĐỌC – VIẾT B2.2"; COURSE_EN="Reading & Writing B2.2"; TIME="40 phút"
ORDER=[16,17,18,19,20,22]   # mã đề 01..06
E=G.E
def code(i): return f"RW22-R-{i:02d}"

def load(n):
    T=importlib.import_module(f't{n}').T
    drop=set(T['drop']); allq=sorted(T['key']); kept=[q for q in allq if q not in drop]
    T['kept']=kept; T['map']={o:i+1 for i,o in enumerate(kept)}; T['inv']={v:k for k,v in T['map'].items()}
    # passage boundaries: passage 0 = questions of groups with p==0
    T['pq']=[[T['map'][q] for q in kept if any(g['p']==pi and q in list(g['qs']) for g in T['groups'])] for pi in (0,1)]
    return T
TESTS=[load(n) for n in ORDER]

def qrange(qs,T):
    ks=[T['map'][q] for q in qs if q in T['map']]
    if not ks: return None
    return f"Question {ks[0]}" if len(ks)==1 else f"Questions {ks[0]}–{ks[-1]}"
def fill(text,T,blank='______'):
    return re.sub(r'\{(\d+)\}',lambda m:f"({T['map'][int(m.group(1))]}) {blank}" if int(m.group(1)) in T['map'] else '',text)
def show_key(T,o):
    return T['key'][o] if not (T['kind'].get(o) in ('opt',)) else T['key'][o]
def kind_map(T):
    d={}
    for g in T['groups']:
        for q in g['qs']: d[q]=g['kind']
    T['kind']=d
for T in TESTS: kind_map(T)
def tf_labels(g): return "YES / NO / NOT GIVEN" if g.get('yn') else "TRUE / FALSE / NOT GIVEN"

# ---------------- docx ----------------
def top_block(d,i,teacher=False):
    if teacher: P(d,'TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN',b=True,sz=9.5,color=RED,align='c',after=3)
    tb=d.add_table(rows=1,cols=2); noborder(tb); widths(tb,[8.2,10]); l,r=tb.rows[0].cells
    cellp(l,'BỘ XÂY DỰNG',sz=10.5,align='c'); cellp(l,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM',b=True,sz=11,align='c',first=False); cellp(l,'KHOA NGOẠI NGỮ',b=True,sz=11,align='c',first=False); cellp(l,'—————',sz=10.5,align='c',first=False)
    cellp(r,('ĐÁP ÁN VÀ GIẢI THÍCH' if teacher else 'ĐỀ KIỂM TRA GIỮA KỲ')+'  –  READING',b=True,sz=16,color=NAVY,align='c'); cellp(r,YEAR,sz=10.5,align='c',first=False); cellp(r,'Hình thức: Đọc hiểu (25 câu)',i=True,sz=10.5,align='c',first=False)
    P(d,'',after=2)
    tb=d.add_table(rows=1,cols=2); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; widths(tb,[12.4,5.8]); l,r=tb.rows[0].cells
    cellp(l,f'Học phần: {COURSE}  ({COURSE_EN})',b=True,sz=11); cellp(l,f'Thời gian làm bài: {TIME}',sz=10.5,first=False)
    shade(r,PINK); cellp(r,'MÃ ĐỀ',b=True,sz=10,align='c'); cellp(r,f'{i:02d}',b=True,sz=22,color=RED,align='c',first=False)

def box_options(d,opts,cols):
    rows=(len(opts)+cols-1)//cols; tb=d.add_table(rows=rows,cols=cols); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for k,(L,txt) in enumerate(opts):
        c=tb.rows[k//cols].cells[k%cols]; shade(c,LBLUE); p=c.paragraphs[0]; R(p,L+'  ',b=True,color=NAVY,sz=10); R(p,txt,sz=9.5); p.paragraph_format.space_after=Pt(1)
    widths(tb,[18.2/cols]*cols)
def box_list(d,opts):  # single column list (headings)
    tb=d.add_table(rows=len(opts),cols=1); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for k,(L,txt) in enumerate(opts):
        c=tb.rows[k].cells[0]; shade(c,LBLUE); p=c.paragraphs[0]; R(p,f'{L}   ',b=True,color=NAVY,sz=10); R(p,txt,sz=10); p.paragraph_format.space_after=Pt(0)
    widths(tb,[18.2])

def passage_docx(d,T,pi,title_no):
    qs=T['pq'][pi]
    part_heading(d,f"PASSAGE {title_no}  (Questions {qs[0]}–{qs[-1]})",f"{len(qs)} questions · about 20 minutes")
    ps=T['passages'][pi]
    P(d,ps['title'],b=True,sz=14,color=NAVY,align='c',before=2,after=4)
    for letter,text in ps['paras']:
        for k,chunk in enumerate(text.split('\n\n')):
            p=P(d,'',after=4); 
            if letter and k==0: R(p,letter+'   ',b=True,color=NAVY)
            elif letter: p.paragraph_format.left_indent=Cm(0.55)
            R(p,chunk,sz=10.5)
    if ps.get('note'): P(d,ps['note'],i=True,sz=9,color=GREY,after=4)
    for g in [g for g in T['groups'] if g['p']==pi]:
        rg=qrange(g['qs'],T)
        if not rg: continue
        p=P(d,'',before=8,after=2,keep=True); R(p,rg,b=True,color=NAVY,sz=11.5)
        P(d,g['instr'],b=True,i=True,sz=10,color=NAVY,after=3,keep=True)
        k=g['kind']
        if k in('heading',): P(d,'List of Headings',b=True,sz=10,after=2,keep=True); box_list(d,g['opts'])
        if k=='match' or (k=='summary' and g.get('opts')):
            box_options(d,g['opts'],cols=(4 if len(g['opts'])>6 else len(g['opts'])) if k=='summary' else min(len(g['opts']),4))
        if k=='summary':
            P(d,fill(g['text'],T),sz=10.5,before=4,after=3)
        elif k=='mcq':
            for o in g['qs']:
                if o not in T['map']: continue
                q,opts=g['items'][o]; p=P(d,'',before=5,after=1,keep=True); R(p,f"{T['map'][o]}.  ",b=True,color=NAVY); R(p,q)
                for L in 'ABCD':
                    pp=P(d,'',after=0,indent=0.7,keep=(L!='D')); R(pp,f"{L}.  ",b=True,color=NAVY); R(pp,opts[L],sz=10.5)
        elif k=='tfng':
            for o in g['qs']:
                if o not in T['map']: continue
                p=P(d,'',before=4,after=1); R(p,f"{T['map'][o]}.  ",b=True,color=NAVY); R(p,g['items'][o]); R(p,'   ______',color=GREY)
        elif k=='completion':
            for o in g['qs']:
                if o not in T['map']: continue
                p=P(d,'',before=4,after=1); R(p,f"{T['map'][o]}.  ",b=True,color=NAVY); R(p,g['items'][o].replace('______','________'))
        elif k in('match','para','heading'):
            for o in g['qs']:
                if o not in T['map']: continue
                p=P(d,'',before=4,after=1); R(p,f"{T['map'][o]}.  ",b=True,color=NAVY); R(p,g['items'][o]+('' if g['items'][o].endswith('.') else ':')+'  '); R(p,'______',color=GREY)

def exam_docx(path,T,i):
    d=base_doc(); footer(d,f"{code(i)}  ·  {COURSE_EN} – Reading  ·  Mã đề {i:02d}")
    top_block(d,i)
    p=P(d,'',before=4,after=1); R(p,'Không sử dụng tài liệu [x]      Được sử dụng tài liệu [   ]      Nộp lại đề thi [x]',sz=10)
    P(d,'Đề thi gồm 25 câu, 2 bài đọc (Passage 1, Passage 2), kèm 01 Phiếu trả lời ở trang cuối. Cán bộ coi thi không giải thích gì thêm.',sz=10,after=1)
    P(d,'Sinh viên ghi MÃ ĐỀ và làm bài trên PHIẾU TRẢ LỜI.',b=True,i=True,sz=10,color=NAVY,after=1)
    rule(d)
    passage_docx(d,T,0,1); d.add_page_break(); passage_docx(d,T,1,2)
    P(d,'— END OF THE TEST —',b=True,color=NAVY,align='c',before=12,after=6)
    tb=d.add_table(rows=1,cols=2); noborder(tb); widths(tb,[9,9])
    for c,role in zip(tb.rows[0].cells,('Trưởng Bộ môn duyệt','Giảng viên ra đề')):
        cellp(c,'TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026',i=True,sz=10,align='c'); cellp(c,role,b=True,sz=10.5,align='c',first=False); cellp(c,'(Ký và ghi rõ họ, tên)',i=True,sz=9.5,align='c',first=False)
        for _ in range(3): cellp(c,'',first=False)
    d.add_page_break()
    P(d,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM  ·  KHOA NGOẠI NGỮ',b=True,sz=10.5,color=NAVY,align='c',after=0)
    P(d,'PHIẾU TRẢ LỜI  /  STUDENT ANSWER SHEET',b=True,sz=17,color=NAVY,align='c',after=0)
    P(d,f'Kiểm tra giữa kỳ · {COURSE} (READING) · {YEAR}',i=True,sz=10,color=GREY,align='c',after=6)
    tb=d.add_table(rows=3,cols=2); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; widths(tb,[13,5.2])
    for r_,txt in zip(tb.rows,['Họ và tên:  ..................................................................','MSSV:  ...................................    Lớp:  ..............................','Phòng thi:  ..........    Số tờ:  ..........    Chữ ký SV:  ....................']): cellp(r_.cells[0],txt,sz=10.5)
    m=tb.cell(0,1).merge(tb.cell(2,1)); shade(m,PINK); m.text=''; cellp(m,'MÃ ĐỀ',b=True,sz=10,align='c'); cellp(m,f'{i:02d}',b=True,sz=26,color=RED,align='c',first=False)
    P(d,'',after=3)
    tb=d.add_table(rows=2,cols=4); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(tb.rows[0].cells,['TỔNG (/10)','Số câu đúng (/25)','GV chấm 1','GV chấm 2']): cellp(c,h,b=True,sz=9.5,align='c'); shade(c,LBLUE)
    for c in tb.rows[1].cells: cellp(c,'')
    tb.rows[1].height=Cm(1.0)
    p=P(d,'',before=8,after=3); R(p,'ANSWERS',b=True,color=NAVY,sz=11.5); R(p,'   — Write ONLY the answer (letter, Roman numeral, word(s) or TRUE / FALSE / NOT GIVEN / YES / NO). Spelling must be correct. 0.4 point per question.',i=True,sz=9.5)
    tb=d.add_table(rows=14,cols=4); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['Câu','Your answer','Câu','Your answer']): cellp(tb.rows[0].cells[j],h,b=True,sz=9.5,align='c'); shade(tb.rows[0].cells[j],LBLUE)
    for r in range(13):
        for blk in range(2):
            n=blk*13+r+1
            if n<=25: cellp(tb.rows[r+1].cells[blk*2],f'{n}.',b=True,sz=10.5,align='c')
            cellp(tb.rows[r+1].cells[blk*2+1],'')
        tb.rows[r+1].height=Cm(0.85)
    widths(tb,[1.2,7.9,1.2,7.9])
    d.save(path)

def key_docx(path,T,i):
    d=base_doc(); footer(d,f"{code(i)}  ·  Đáp án (lưu hành nội bộ)")
    top_block(d,i,teacher=True); P(d,'',after=2)
    p=P(d,'',after=1); R(p,'Nguồn: ',b=True,sz=9.5); R(p,T['src']+'. Thư mục Google Drive "IELTS Reading Actual Tests 1" của giảng viên.',sz=9.5)
    p=P(d,'',after=3); R(p,'Cấu trúc: ',b=True,sz=9.5); R(p,f"25 câu = Passage 1 (câu 1–{T['pq'][0][-1]}, gốc Q14–25) + Passage 2 (câu {T['pq'][1][0]}–25, gốc Q27–39). Đã bỏ câu cuối của mỗi passage (gốc Q26 và Q40). Bảng đối chiếu số câu nằm trong cột \"Câu gốc\".",sz=9.5)
    part_heading(d,'ĐÁP ÁN NHANH','25 câu · 0.4 điểm/câu')
    n=len(T['kept']); rows=13
    tb=d.add_table(rows=rows+1,cols=6); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['Câu','Câu gốc','Đáp án']*2): cellp(tb.rows[0].cells[j],h,b=True,sz=9,align='c'); shade(tb.rows[0].cells[j],LBLUE)
    for r in range(rows):
        for blk in range(2):
            k=blk*rows+r+1
            if k<=n:
                o=T['inv'][k]; cellp(tb.rows[r+1].cells[blk*3],f'{k}.',b=True,sz=9.5,align='c'); cellp(tb.rows[r+1].cells[blk*3+1],f'Q{o}',sz=9,color=GREY,align='c'); cellp(tb.rows[r+1].cells[blk*3+2],T['key'][o],b=True,sz=9.5,color=RED)
            else:
                for c in range(3): cellp(tb.rows[r+1].cells[blk*3+c],'')
    widths(tb,[1.0,1.5,6.6]*2)
    P(d,'Thang điểm: 25 câu × 0,4 = 10 điểm. Câu điền từ: đúng nội dung và chính tả, đúng giới hạn từ; không trừ lỗi viết hoa. Dấu nháy và gạch ngang được chấp nhận. Mọi cách viết ghi sau dấu "/" đều được chấp nhận.',sz=10,before=4,after=2)
    fixes=[(o,T['fixnote'][o]) for o in sorted(T.get('fixnote',{})) if o in T['map']]
    if fixes:
        part_heading(d,'LƯU Ý KHI KIỂM TRA ĐÁP ÁN NGUỒN','')
        for o,t in fixes:
            p=P(d,'',after=2); R(p,f"Câu {T['map'][o]} (gốc Q{o}): ",b=True,sz=10,color=RED); R(p,t,sz=10)
    for pi in (0,1):
        part_heading(d,f"GIẢI THÍCH CHI TIẾT – PASSAGE {pi+1}",T['passages'][pi]['title'])
        for g in [g for g in T['groups'] if g['p']==pi]:
            for o in g['qs']:
                if o not in T['map']: continue
                k=T['map'][o]
                p=P(d,'',before=6,after=1,keep=True); R(p,f'Câu {k}  ',b=True,color=NAVY,sz=11); R(p,f'(gốc Q{o}, {gname(g)})',i=True,sz=9.5,color=GREY); R(p,'    Đáp án: ',b=True)
                ans=T['key'][o]
                if g['kind']=='mcq': ans=f"{ans}. {g['items'][o][1][ans]}"
                elif g['kind'] in('match',) : ans=f"{ans} – {dict(g['opts'])[ans]}"
                elif g['kind']=='heading': ans=f"{ans} – {dict(g['opts'])[ans]}"
                elif g['kind']=='summary' and g.get('opts'): ans=f"{ans} – {dict(g['opts'])[ans]}"
                R(p,ans,b=True,color=RED)
                if g['kind']=='mcq': P(d,g['items'][o][0],i=True,sz=10,indent=0.6,after=1,keep=True)
                elif g['kind'] in('tfng','completion','match','para','heading'): P(d,(g['items'][o] if isinstance(g['items'][o],str) else ''),i=True,sz=10,indent=0.6,after=1,keep=True)
                pp=P(d,'',indent=0.6,after=2); R(pp,'Giải thích: ',b=True,sz=10); R(pp,T['why'][o],sz=10)
    d.save(path)
def gname(g):
    return {'match':'Matching','para':'Matching information','heading':'Matching headings','mcq':'Multiple choice','tfng':'Yes/No/Not Given' if g.get('yn') else 'True/False/Not Given','summary':'Summary completion','completion':'Sentence completion'}[g['kind']]

def quick_docx(path):
    from docx.enum.section import WD_ORIENT
    d=base_doc(); sec=d.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=sec.page_height,sec.page_width
    P(d,'TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN',b=True,sz=9.5,color=RED,align='c',after=2)
    P(d,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM  ·  KHOA NGOẠI NGỮ',b=True,sz=10.5,color=NAVY,align='c',after=0)
    P(d,'BẢNG ĐÁP ÁN NHANH — 6 MÃ ĐỀ READING',b=True,sz=17,color=NAVY,align='c',after=0)
    P(d,f'Kiểm tra giữa kỳ · {COURSE} (READING) · 25 câu · mỗi câu 0,4đ',i=True,sz=10,color=GREY,align='c',after=6)
    for i,T in enumerate(TESTS,1):
        P(d,f"Mã đề {i:02d} — Passage 1: {T['passages'][0]['title']} · Passage 2: {T['passages'][1]['title']}",b=True,sz=10,color=NAVY,before=4,after=1)
        tb=d.add_table(rows=2,cols=26); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
        cellp(tb.rows[0].cells[0],'Câu',b=True,sz=8,align='c'); shade(tb.rows[0].cells[0],LBLUE); cellp(tb.rows[1].cells[0],'Đ/á',b=True,sz=8,align='c')
        for k in range(1,26):
            cellp(tb.rows[0].cells[k],str(k),b=True,sz=7.5,align='c'); shade(tb.rows[0].cells[k],LBLUE)
            a=T['key'][T['inv'][k]]; cellp(tb.rows[1].cells[k],a.split(' / ')[0].replace(' (to choose)','').replace('(s)',''),sz=6.5,align='c')
        widths(tb,[1.0]+[0.98]*25)
    P(d,'Chi tiết các cách viết được chấp nhận và giải thích: xem file đáp án từng mã đề.',sz=9.5,before=6)
    d.save(path)

# ---------------- html ----------------
CSS=G.CSS+"<style>.pt{font-size:14pt;font-weight:bold;text-align:center;color:#"+NAVY+";margin:4px 0}.pp{margin:3px 0 6px;font-size:10.5pt;text-align:justify}.pl{font-weight:bold;color:#"+NAVY+"}.bxo td{border:1px solid #444;background:#"+LBLUE+";padding:2px 5px;font-size:9.5pt;width:25%}.bxl td{border:1px solid #444;background:#"+LBLUE+";padding:2px 6px;font-size:10pt}.opt{margin-left:.7cm}.gh{margin-top:9px;font-weight:bold;color:#"+NAVY+";font-size:11.5pt;break-after:avoid}.gi{font-weight:bold;font-style:italic;color:#"+NAVY+";font-size:10pt;margin:1px 0 4px;break-after:avoid}</style>"
def htop(i,teacher=False):
    h=("<p class='c b red' style='font-size:9.5pt;margin:0 0 4px'>TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN</p>" if teacher else "")
    ttl=('ĐÁP ÁN VÀ GIẢI THÍCH' if teacher else 'ĐỀ KIỂM TRA GIỮA KỲ')+'  –  READING'
    return h+(f"<table><tr><td width=42% class=c>BỘ XÂY DỰNG<br><b>HỌC VIỆN HÀNG KHÔNG VIỆT NAM</b><br><b>KHOA NGOẠI NGỮ</b><br>—————</td><td class=c><div class='b nav' style='font-size:16pt'>{ttl}</div>{YEAR}<br><i>Hình thức: Đọc hiểu (25 câu)</i></td></tr></table><div style='height:6px'></div>"
        f"<table class=bx><tr><td width=68%><b style='font-size:11pt'>Học phần: {COURSE}  ({COURSE_EN})</b><br>Thời gian làm bài: {TIME}</td><td class='c pk'><b style='font-size:10pt'>MÃ ĐỀ</b><br><b class=red style='font-size:22pt;line-height:1.1'>{i:02d}</b></td></tr></table>")
def passage_html(T,pi,no):
    qs=T['pq'][pi]; ps=T['passages'][pi]
    h=f"<div class=ph>PASSAGE {no}  (Questions {qs[0]}–{qs[-1]}) <span>&nbsp; {len(qs)} questions · about 20 minutes</span></div><div class=pt>{E(ps['title'])}</div>"
    for letter,text in ps['paras']:
        for k,chunk in enumerate(text.split('\n\n')):
            lab=f"<span class=pl>{letter}</span>&nbsp;&nbsp; " if letter and k==0 else ''
            st=' style="margin-left:.55cm"' if (letter and k>0) else ''
            h+=f"<p class=pp{st}>{lab}{E(chunk)}</p>"
    if ps.get('note'): h+=f"<p class='i gr' style='font-size:9pt'>{E(ps['note'])}</p>"
    for g in [g for g in T['groups'] if g['p']==pi]:
        rg=qrange(g['qs'],T)
        if not rg: continue
        h+=f"<div class=gh>{rg}</div><div class=gi>{E(g['instr'])}</div>"; k=g['kind']
        if k=='heading': h+="<p class=b style='margin:0 0 2px;font-size:10pt'>List of Headings</p><table class=bxl>"+"".join(f"<tr><td><b class=nav>{L}</b> &nbsp; {E(t)}</td></tr>" for L,t in g['opts'])+"</table>"
        if k=='match' or (k=='summary' and g.get('opts')):
            cols=4 if len(g['opts'])>4 or k=='summary' else len(g['opts']); h+="<table class=bxo>"
            for r in range(0,len(g['opts']),cols): h+="<tr>"+"".join(f"<td><b class=nav>{L}</b> {E(t)}</td>" for L,t in g['opts'][r:r+cols])+"</tr>"
            h+="</table>"
        if k=='summary': h+=f"<p class=pp style='margin-top:5px'>{E(fill(g['text'],T))}</p>"
        elif k=='mcq':
            for o in g['qs']:
                if o not in T['map']: continue
                q,opts=g['items'][o]; h+=f"<div class=q><div><b class=nav>{T['map'][o]}.</b>&nbsp; {E(q)}</div>"+"".join(f"<div class=opt><b class=nav>{L}.</b>&nbsp; {E(opts[L])}</div>" for L in 'ABCD')+"</div>"
        elif k in('tfng','completion','match','para','heading'):
            for o in g['qs']:
                if o not in T['map']: continue
                tail=' <span class=gr>______</span>' if k!='completion' else ''
                h+=f"<div class=q style='margin-top:4px'><b class=nav>{T['map'][o]}.</b>&nbsp; {E(g['items'][o]+(':' if (k in ('match','para','heading') and not g['items'][o].endswith('.')) else ''))}{tail}</div>"
    return h
def exam_html(T,i):
    h=f"<html><head><meta charset=utf-8>{CSS}</head><body>"+htop(i)
    h+=("<p style='margin:8px 0 2px;font-size:10pt'>Không sử dụng tài liệu [x] &nbsp;&nbsp;&nbsp; Được sử dụng tài liệu [&nbsp;&nbsp;&nbsp;] &nbsp;&nbsp;&nbsp; Nộp lại đề thi [x]</p><p style='margin:0 0 2px;font-size:10pt'>Đề thi gồm 25 câu, 2 bài đọc (Passage 1, Passage 2), kèm 01 Phiếu trả lời ở trang cuối. Cán bộ coi thi không giải thích gì thêm.</p><p class='b i nav' style='margin:0;font-size:10pt'>Sinh viên ghi MÃ ĐỀ và làm bài trên PHIẾU TRẢ LỜI.</p><div class=rule></div>")
    h+=passage_html(T,0,1)+"<div style='break-before:page'></div>"+passage_html(T,1,2)
    h+="<p class='c b nav' style='margin:14px 0 6px'>— END OF THE TEST —</p><table style='break-inside:avoid'><tr>"+"".join(f"<td class=c style='font-size:10pt'><i>TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026</i><br><b>{r}</b><br><i style='font-size:9.5pt'>(Ký và ghi rõ họ, tên)</i><div style='height:50px'></div></td>" for r in ('Trưởng Bộ môn duyệt','Giảng viên ra đề'))+"</tr></table>"
    h+="<div style='break-before:page'>"
    h+=(f"<p class='c b nav' style='margin:0'>HỌC VIỆN HÀNG KHÔNG VIỆT NAM &nbsp;·&nbsp; KHOA NGOẠI NGỮ</p><p class='c b nav' style='margin:0;font-size:17pt'>PHIẾU TRẢ LỜI &nbsp;/&nbsp; STUDENT ANSWER SHEET</p><p class='c i gr' style='margin:0 0 8px;font-size:10pt'>Kiểm tra giữa kỳ · {COURSE} (READING) · {YEAR}</p>"
        f"<table class=bx><tr><td width=72%>Họ và tên: ..................................................................</td><td rowspan=3 class='c pk'><b style='font-size:10pt'>MÃ ĐỀ</b><br><b class=red style='font-size:26pt'>{i:02d}</b></td></tr><tr><td>MSSV: ................................... &nbsp;&nbsp; Lớp: ..............................</td></tr><tr><td>Phòng thi: .......... &nbsp;&nbsp; Số tờ: .......... &nbsp;&nbsp; Chữ ký SV: ....................</td></tr></table><div style='height:8px'></div>"
        f"<table class='bx c'><tr style='background:#{LBLUE}'><th>TỔNG (/10)</th><th>Số câu đúng (/25)</th><th>GV chấm 1</th><th>GV chấm 2</th></tr><tr style='height:32px'><td></td><td></td><td></td><td></td></tr></table>")
    h+="<p style='margin:12px 0 4px'><b class=nav style='font-size:11.5pt'>ANSWERS</b> <i style='font-size:9.5pt'>— Write ONLY the answer (letter, Roman numeral, word(s) or TRUE / FALSE / NOT GIVEN / YES / NO). Spelling must be correct. 0.4 point per question.</i></p><table class=lt><tr><th>Câu</th><th>Your answer</th><th>Câu</th><th>Your answer</th></tr>"
    for r in range(13): h+=f"<tr style='height:30px'><td width=7%><b>{r+1}.</b></td><td></td><td width=7%><b>{r+14 if r+14<=25 else ''}{'.' if r+14<=25 else ''}</b></td><td></td></tr>"
    return h+"</table></div></body></html>"
def key_html(T,i):
    h=f"<html><head><meta charset=utf-8>{CSS}</head><body>"+htop(i,True)
    h+=f"<p style='margin:6px 0 1px;font-size:9.5pt'><b>Nguồn:</b> {E(T['src'])}. Thư mục Google Drive \"IELTS Reading Actual Tests 1\" của giảng viên.</p><p style='margin:0 0 4px;font-size:9.5pt'><b>Cấu trúc:</b> 25 câu = Passage 1 (câu 1–{T['pq'][0][-1]}, gốc Q14–25) + Passage 2 (câu {T['pq'][1][0]}–25, gốc Q27–39). Đã bỏ câu cuối của mỗi passage (gốc Q26 và Q40).</p>"
    h+="<div class=ph>ĐÁP ÁN NHANH <span>&nbsp; 25 câu · 0.4 điểm/câu</span></div><table class=lt><tr>"+"".join("<th>Câu</th><th>Gốc</th><th>Đáp án</th>" for _ in range(2))+"</tr>"
    n=len(T['kept'])
    for r in range(13):
        h+="<tr>"
        for blk in range(2):
            k=blk*13+r+1
            h+=(f"<td><b>{k}.</b></td><td class=gr>Q{T['inv'][k]}</td><td class='b red' style='text-align:left'>{E(T['key'][T['inv'][k]])}</td>" if k<=n else "<td></td><td></td><td></td>")
        h+="</tr>"
    h+="</table><p style='font-size:10pt;margin:5px 0'><b>Thang điểm:</b> 25 câu × 0,4 = 10 điểm. Câu điền từ: đúng nội dung và chính tả, đúng giới hạn từ; không trừ lỗi viết hoa. Mọi cách viết ghi sau dấu \"/\" đều được chấp nhận.</p>"
    fixes=[(o,T['fixnote'][o]) for o in sorted(T.get('fixnote',{})) if o in T['map']]
    if fixes:
        h+="<div class=ph>LƯU Ý KHI KIỂM TRA ĐÁP ÁN NGUỒN</div>"+"".join(f"<p style='font-size:10pt;margin:2px 0'><b class=red>Câu {T['map'][o]} (gốc Q{o}):</b> {E(t)}</p>" for o,t in fixes)
    for pi in (0,1):
        h+=f"<div class=ph>GIẢI THÍCH CHI TIẾT – PASSAGE {pi+1} <span>&nbsp; {E(T['passages'][pi]['title'])}</span></div>"
        for g in [g for g in T['groups'] if g['p']==pi]:
            for o in g['qs']:
                if o not in T['map']: continue
                ans=T['key'][o]
                if g['kind']=='mcq': ans=f"{ans}. {g['items'][o][1][ans]}"
                elif g['kind'] in('match','heading') or (g['kind']=='summary' and g.get('opts')): ans=f"{ans} – {dict(g['opts'])[ans]}"
                stem=(g['items'][o][0] if g['kind']=='mcq' else (g['items'][o] if g['kind'] in('tfng','completion','match','para','heading') else ''))
                h+=(f"<div class=blk><b class=nav style='font-size:11pt'>Câu {T['map'][o]}</b> <i class=gr style='font-size:9.5pt'>(gốc Q{o}, {gname(g)})</i> &nbsp;<b>Đáp án: <span class=red>{E(ans)}</span></b>"
                    +(f"<div class=k style='font-size:10pt'><i>{E(stem)}</i></div>" if stem else '')+f"<div class=k style='font-size:10pt'><b>Giải thích:</b> {E(T['why'][o])}</div></div>")
    return h+"</body></html>"
def quick_html():
    h=f"<html><head><meta charset=utf-8>{CSS.replace('size:Letter','size:Letter landscape')}</head><body>"
    h+=(f"<p class='c b red' style='font-size:9.5pt;margin:0'>TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN</p><p class='c b nav' style='margin:0'>HỌC VIỆN HÀNG KHÔNG VIỆT NAM &nbsp;·&nbsp; KHOA NGOẠI NGỮ</p><p class='c b nav' style='margin:0;font-size:17pt'>BẢNG ĐÁP ÁN NHANH — 6 MÃ ĐỀ READING</p><p class='c i gr' style='margin:0 0 6px;font-size:10pt'>Kiểm tra giữa kỳ · {COURSE} (READING) · 25 câu · mỗi câu 0,4đ</p>")
    for i,T in enumerate(TESTS,1):
        h+=f"<p class='b nav' style='margin:6px 0 1px;font-size:10pt'>Mã đề {i:02d} — Passage 1: {E(T['passages'][0]['title'])} · Passage 2: {E(T['passages'][1]['title'])}</p><table class=lt><tr><th>Câu</th>"+"".join(f"<th>{k}</th>" for k in range(1,26))+"</tr><tr><td><b>Đ/á</b></td>"+"".join(f"<td style='font-size:7pt'>{E(T['key'][T['inv'][k]].split(' / ')[0].replace(' (to choose)',''))}</td>" for k in range(1,26))+"</tr></table>"
    return h+"<p style='font-size:9.5pt'>Chi tiết các cách viết được chấp nhận và giải thích: xem file đáp án từng mã đề.</p></body></html>"

def main():
    hout=os.environ.get('HTML_OUT')
    dt=os.path.join(OUT,'De_thi_va_Phieu_tra_loi'); da=os.path.join(OUT,'Dap_an'); os.makedirs(dt,exist_ok=True); os.makedirs(da,exist_ok=True)
    for i,T in enumerate(TESTS,1):
        nm=code(i)
        exam_docx(os.path.join(dt,f'{nm} - DE THI.docx'),T,i); key_docx(os.path.join(da,f'{nm} - DAP AN.docx'),T,i)
        if hout:
            for sub,fn,hh in (('De_thi_va_Phieu_tra_loi',f'{nm} - DE THI',exam_html(T,i)),('Dap_an',f'{nm} - DAP AN',key_html(T,i))):
                pd=os.path.join(hout,sub); os.makedirs(pd,exist_ok=True); open(os.path.join(pd,fn+'.html'),'w',encoding='utf-8').write(hh)
    quick_docx(os.path.join(da,'00 - BANG DAP AN NHANH 6 MA DE.docx'))
    if hout: open(os.path.join(hout,'Dap_an','00 - BANG DAP AN NHANH 6 MA DE.html'),'w',encoding='utf-8').write(quick_html())
if __name__=="__main__": main()
