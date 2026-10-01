# -*- coding: utf-8 -*-
"""Tạo đề Listening giữa kỳ LS1 (docx + html cho pdf) từ 12 test nguồn.  python3 build_listening.py <thư mục LS1 folder> <thư mục audio đã tạo>"""
import os,sys,re,json,html as _h
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
_argv=sys.argv; sys.argv=[_argv[0]]
sys.path.insert(0,os.path.join(HERE,'..','_build')); import build as G
sys.argv=_argv
import data,explain
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R,P,cellp,shade,shade_par,widths,noborder,rule,footer,part_heading,base_doc=G.R,G.P,G.cellp,G.shade,G.shade_par,G.widths,G.noborder,G.rule,G.footer,G.part_heading,G.base_doc
NAVY,PINK,LBLUE,RED,GREY=G.NAVY,G.PINK,G.LBLUE,G.RED,G.GREY
OUT=sys.argv[1] if len(sys.argv)>1 else os.path.join(HERE,'..','LS1-MIDTERM TEST - LISTENING')
AUD=sys.argv[2] if len(sys.argv)>2 else os.path.join(OUT,'Audio')
YEAR=G.YEAR; COURSE="NGHE – NÓI 1"; COURSE_EN="Listening and Speaking 1"
DUR=json.load(open(os.path.join(AUD,'durations.json'))) if os.path.exists(os.path.join(AUD,'durations.json')) else {}
TIME="10 phút (gồm 2 file ghi âm, mỗi file chỉ nghe 1 lần)"
LISTING_URL="https://youpass.vn/luyen-thi/ielts/listening?quiz_type=mocktest&status=unfinished&mock_source=youpass_collect"
def code(n): return f"LS1-L-{n:02d}"
def audio_href(n,part): return f"../Audio/{code(n)}_Part{part}.mp3"
def fmt_dur(sec): m=int(sec//60); s=int(round(sec-60*m)); return f"{m}:{s:02d}"

# ---------- docx helpers
def add_link(p,url,text,sz=11,bold=True):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rPr.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u)
    if bold: rPr.append(OxmlElement('w:b'))
    s=OxmlElement('w:sz'); s.set(qn('w:val'),str(int(sz*2))); rPr.append(s)
    r.append(rPr); t=OxmlElement('w:t'); t.text=text; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r); p._p.append(h)

def top_block(d,t,teacher=False):
    n=t['n']
    if teacher: P(d,'TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN',b=True,sz=9.5,color=RED,align='c',after=3)
    tb=d.add_table(rows=1,cols=2); noborder(tb); widths(tb,[8.2,10])
    l,r=tb.rows[0].cells
    cellp(l,'BỘ XÂY DỰNG',sz=10.5,align='c'); cellp(l,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM',b=True,sz=11,align='c',first=False); cellp(l,'KHOA NGOẠI NGỮ',b=True,sz=11,align='c',first=False); cellp(l,'—————',sz=10.5,align='c',first=False)
    cellp(r,('ĐÁP ÁN VÀ GIẢI THÍCH' if teacher else 'ĐỀ KIỂM TRA GIỮA KỲ')+'  –  LISTENING',b=True,sz=16,color=NAVY,align='c'); cellp(r,YEAR,sz=10.5,align='c',first=False)
    cellp(r,'Hình thức: Nghe (2 file ghi âm)',i=True,sz=10.5,align='c',first=False)
    P(d,'',after=2)
    tb=d.add_table(rows=1,cols=2); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; widths(tb,[12.4,5.8])
    l,r=tb.rows[0].cells
    cellp(l,f'Học phần: {COURSE}  ({COURSE_EN})',b=True,sz=11); cellp(l,f'Thời gian làm bài: {TIME}',sz=10.5,first=False)
    shade(r,PINK); cellp(r,'MÃ ĐỀ',b=True,sz=10,align='c'); cellp(r,f'{n:02d}',b=True,sz=22,color=RED,align='c',first=False)

def exam_docx(path,t):
    n=t['n']; d=base_doc(); footer(d,f"{code(n)}  ·  {COURSE_EN} – Listening  ·  Mã đề {n:02d}")
    top_block(d,t)
    p=P(d,'',before=4,after=1); R(p,'Không sử dụng tài liệu [x]      Được sử dụng tài liệu [   ]      Nộp lại đề thi [x]',sz=10)
    P(d,'Đề thi gồm 20 câu, nghe 2 file ghi âm (Part 1, Part 2), kèm 01 Phiếu trả lời ở trang cuối. Mỗi file chỉ được nghe MỘT LẦN. Cán bộ coi thi không giải thích gì thêm.',sz=10,after=1)
    P(d,'Sinh viên ghi MÃ ĐỀ và làm bài trên PHIẾU TRẢ LỜI.',b=True,i=True,sz=10,color=NAVY,after=3)
    # recordings box
    tb=d.add_table(rows=1,cols=1); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; widths(tb,[18.2]); c=tb.rows[0].cells[0]; shade(c,LBLUE)
    cellp(c,'RECORDINGS  /  FILE GHI ÂM  (nhấn vào liên kết để mở file nghe)',b=True,sz=10.5,color=NAVY)
    for part,qs in ((1,'Questions 1–10'),(2,'Questions 11–20')):
        pp=c.add_paragraph(); pp.paragraph_format.space_after=Pt(1); R(pp,f'▶  Part {part} ({qs}):  ',b=True,sz=10.5)
        add_link(pp,audio_href(n,part),f'{code(n)}_Part{part}.mp3',sz=10.5)
        R(pp,f'   (≈ {fmt_dur(DUR[str(n)][part-1]) if DUR else "2:30"} phút)',i=True,sz=9.5,color=GREY)
    pp=c.add_paragraph(); R(pp,'Lưu ý: mở file từ thư mục "Audio" nằm cạnh thư mục đề thi; không đổi tên hoặc di chuyển file.',i=True,sz=9,color=GREY)
    P(d,'',after=2); rule(d)
    part_heading(d,'PART 1.  FORM COMPLETION  (Questions 1–10)','5 points · 0.5 point each')
    P(d,t['p1_directions']+' Write your answers on the ANSWER SHEET.',b=True,i=True,sz=10.5,color=NAVY,after=3)
    tb=d.add_table(rows=1+len(t['p1']),cols=2); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tb.rows[0].cells[0].merge(tb.rows[0].cells[1]); shade(hdr,NAVY); cellp(hdr,t['p1_title'],b=True,sz=11,color='FFFFFF',align='c')
    for r_,p_ in zip(tb.rows[1:],t['p1']):
        cellp(r_.cells[0],f"{p_['n']}.",b=True,sz=10.5,color=NAVY,align='c'); cellp(r_.cells[1],p_['label'],sz=10.5)
        r_.height=Cm(0.75)
    widths(tb,[1.4,16.8])
    d.add_page_break()
    part_heading(d,'PART 2.  MULTIPLE CHOICE & MATCHING  (Questions 11–20)','5 points · 0.5 point each')
    P(d,t['p2_title'],b=True,sz=11,color=NAVY,after=2)
    P(d,'Questions 11–15.  Choose the correct letter, A, B or C.',b=True,i=True,sz=10.5,color=NAVY,after=2)
    for q in t['mc']:
        p=P(d,'',before=5,after=1,keep=True); R(p,f"{q['n']}.  ",b=True,color=NAVY); R(p,q['q'])
        for L in 'ABC':
            pp=P(d,'',after=0,indent=0.7,keep=(L!='C')); R(pp,f"{L}.  ",b=True,color=NAVY); R(pp,q['opts'][L])
    P(d,'',after=2)
    P(d,t['match_directions']+'  Write the correct letter, A–E, on the ANSWER SHEET.',b=True,i=True,sz=10.5,color=NAVY,after=3,before=6)
    tb=d.add_table(rows=1,cols=5); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,(L,txt) in zip(tb.rows[0].cells,t['match_opts']):
        shade(c,LBLUE); cellp(c,L,b=True,sz=11,color=NAVY,align='c'); cellp(c,txt,sz=10,align='c',first=False)
    widths(tb,[3.64]*5)
    for it in t['match_items']:
        p=P(d,'',before=5,after=1); R(p,f"{it['n']}.  ",b=True,color=NAVY); R(p,it['label']+':  '); R(p,'______',color=GREY)
    P(d,'— END OF THE TEST —',b=True,color=NAVY,align='c',before=12,after=6)
    tb=d.add_table(rows=1,cols=2); noborder(tb); widths(tb,[9,9])
    for c,role in zip(tb.rows[0].cells,('Trưởng Bộ môn duyệt','Giảng viên ra đề')):
        cellp(c,'TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026',i=True,sz=10,align='c'); cellp(c,role,b=True,sz=10.5,align='c',first=False); cellp(c,'(Ký và ghi rõ họ, tên)',i=True,sz=9.5,align='c',first=False)
        for _ in range(3): cellp(c,'',first=False)
    d.add_page_break()
    # answer sheet
    P(d,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM  ·  KHOA NGOẠI NGỮ',b=True,sz=10.5,color=NAVY,align='c',after=0)
    P(d,'PHIẾU TRẢ LỜI  /  STUDENT ANSWER SHEET',b=True,sz=17,color=NAVY,align='c',after=0)
    P(d,f'Kiểm tra giữa kỳ · {COURSE} (LISTENING) · {YEAR}',i=True,sz=10,color=GREY,align='c',after=6)
    tb=d.add_table(rows=3,cols=2); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; widths(tb,[13,5.2])
    for r_,txt in zip(tb.rows,['Họ và tên:  ..................................................................','MSSV:  ...................................    Lớp:  ..............................','Phòng thi:  ..........    Số tờ:  ..........    Chữ ký SV:  ....................']): cellp(r_.cells[0],txt,sz=10.5)
    m=tb.cell(0,1).merge(tb.cell(2,1)); shade(m,PINK); m.text=''; cellp(m,'MÃ ĐỀ',b=True,sz=10,align='c'); cellp(m,f'{n:02d}',b=True,sz=26,color=RED,align='c',first=False)
    P(d,'',after=3)
    tb=d.add_table(rows=2,cols=5); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(tb.rows[0].cells,['Part 1 (/5)','Part 2 (/5)','TỔNG (/10)','GV chấm 1','GV chấm 2']): cellp(c,h,b=True,sz=9.5,align='c'); shade(c,LBLUE)
    for c,h in zip(tb.rows[1].cells,['......... / 5','......... / 5','','','']): cellp(c,h,sz=10,align='c')
    tb.rows[1].height=Cm(1.0)
    p=P(d,'',before=8,after=3); R(p,'PART 1.',b=True,color=NAVY,sz=11.5); R(p,'   — Write NO MORE THAN THREE WORDS AND/OR A NUMBER. Spelling must be correct.',i=True,sz=9.5)
    tb=d.add_table(rows=6,cols=4); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['Câu','Your answer','Câu','Your answer']): cellp(tb.rows[0].cells[j],h,b=True,sz=9.5,align='c'); shade(tb.rows[0].cells[j],LBLUE)
    for r in range(5):
        for blk in range(2):
            cellp(tb.rows[r+1].cells[blk*2],f'{blk*5+r+1}.',b=True,sz=10.5,align='c'); cellp(tb.rows[r+1].cells[blk*2+1],'')
        tb.rows[r+1].height=Cm(0.95)
    widths(tb,[1.2,7.9,1.2,7.9])
    p=P(d,'',before=8,after=3); R(p,'PART 2.',b=True,color=NAVY,sz=11.5); R(p,'   — Questions 11–15: circle ONE letter.  Questions 16–20: circle ONE letter A–E.',i=True,sz=9.5)
    tb=d.add_table(rows=6,cols=10); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    heads=['Câu','A','B','C','Câu','A','B','C','D','E']
    for j,h in enumerate(heads): cellp(tb.rows[0].cells[j],h,b=True,sz=9.5,align='c'); shade(tb.rows[0].cells[j],LBLUE)
    for r in range(5):
        cellp(tb.rows[r+1].cells[0],f'{11+r}.',b=True,sz=10.5,align='c')
        for k,L in enumerate('ABC'): cellp(tb.rows[r+1].cells[1+k],L,sz=10,color=GREY,align='c')
        cellp(tb.rows[r+1].cells[4],f'{16+r}.',b=True,sz=10.5,align='c')
        for k,L in enumerate('ABCDE'): cellp(tb.rows[r+1].cells[5+k],L,sz=10,color=GREY,align='c')
    widths(tb,[1.3,1.2,1.2,1.2,1.3,1.2,1.2,1.2,1.2,1.2])
    d.save(path)

# ---------- answer helpers
W={'1':'one','2':'two','3':'three','4':'four','5':'five','6':'six','7':'seven','8':'eight','9':'nine','10':'ten','12':'twelve','34':'thirty-four','38':'thirty-eight','42':'forty-two','47':'forty-seven','55':'fifty-five','58':'fifty-eight','72':'seventy-two','120':'one hundred and twenty','150':'one hundred and fifty','300':'three hundred','310':'three hundred and ten','540':'five hundred and forty','865':'eight hundred and sixty-five','45':'forty-five'}
ACCEPT={ (1,6):'120 / £120 / one hundred (and) twenty',(3,2):'25 April / April 25 / 25th April / April 25th',(3,7):'7:30 / 7.30 / seven thirty / half past seven',(3,4):'45 minutes / 45 / forty-five minutes',(3,3):'12 / twelve',(3,9):'150 / £150',(8,9):'1 June / June 1 / 1st June / June 1st',(8,4):'greenmail.com / greenmail dot com',(8,5):'2 / two / 2 years / two years',(9,3):'2 / two / second',(9,5):'3 / three',(9,7):'12 / £12 / twelve',(10,6):'18 July / July 18 / 18th July',(10,7):'9:30 / 9.30 / nine thirty / half past nine',(11,3):'2 / two / both',(11,9):'55 / $55',(11,4):'yes / hire a boat',(12,2):'10 km / 10 kilometres / ten km',(12,8):'58 minutes / 58 / 58 mins',(12,9):'300 / £300',(7,8):'310 / $310',(5,8):'42 / $42',(5,1):'10 March / March 10 / 10th March',(5,7):'3 / three',(6,5):'38 / £38',(6,6):'540 / five hundred (and) forty',(8,7):'72 / £72',(4,10):'865 / £865',(4,7):'20 minutes / 20',(2,10):'020 7946 3812 (có/không khoảng trắng)',(11,7):'07700 431226 (có/không khoảng trắng)',(12,6):'07822 519640 (có/không khoảng trắng)',(2,3):'prices / the prices',(10,6):'18 July / July 18 / 18th July',(12,3):'34 / thirty-four',(1,3):'P784215 (viết liền, đúng chữ cái và số)'}
def accept_for(t,q):
    a=t['key'][q]
    if (t['n'],q) in ACCEPT: return ACCEPT[(t['n'],q)]
    if re.fullmatch(r'\d+',a) and a in W: return f'{a} / {W[a]}'
    return a
def p1_type(t,p):
    a=t['key'][p['n']]
    line=[l for l in t['script1'] if f"[Q{p['n']}:" in l][0]
    spelled=bool(re.search(r'\b[A-Z](?:-[A-Z]){2,}\b',line))
    if spelled: return 'Tên riêng được đánh vần từng chữ cái – cần viết đúng chính tả (sai 1 chữ cái = 0 điểm).'
    if re.search(r'[A-Z]{1,2}\d|\d{3} \d',a) or re.fullmatch(r'[A-Z]{1,2}\d+',a) or re.search(r'\d{4,}',a.replace(' ','')): return 'Mã số/số điện thoại được đọc từng chữ số hoặc từng chữ cái – ghi đủ và đúng thứ tự.'
    if re.search(r'\b(January|February|March|April|May|June|July|August|September|October|November|December)\b',a) or re.search(r'Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday',a): return 'Ngày/thứ – nghe số thứ tự (twenty-fifth = 25) và tên tháng/thứ.'
    if re.search(r'\d',a): return 'Số được đọc thành lời (vd: forty-two = 42) – đổi sang chữ số; chú ý đơn vị đi kèm.'
    return 'Từ khóa nghe trực tiếp trong câu trả lời; chú ý chính tả.'
def p1_evidence(t,n):
    L=t['script1']; k=[i for i,l in enumerate(L) if f"[Q{n}:" in l][0]
    clean=lambda s:re.sub(r'\s*\[Q\d+: [^\]]*\]','',s)
    return (clean(L[k-1]) if k>0 and not re.search(r'\[Q\d+:',L[k-1]) else ''), clean(L[k])

def key_docx(path,t):
    n=t['n']; d=base_doc(); footer(d,f"{code(n)}  ·  Đáp án (lưu hành nội bộ)")
    top_block(d,t,teacher=True); P(d,'',after=2)
    p=P(d,'',after=1); R(p,'Test '+str(n)+': ',b=True,color=NAVY); R(p,t['title'],b=True)
    p=P(d,'',after=1); R(p,'Nguồn tham chiếu: ',b=True,sz=9.5); R(p,f"{t['anchor']} – ",sz=9.5); add_link(p,t['url'],t['url'],sz=9.5,bold=False)
    P(d,'Lưu ý: file nguồn (YouPass_Listening_Parts_1_2_Teacher_Edition_12_Tests) ghi rõ câu hỏi, đáp án và transcript là bản VIẾT LẠI theo chủ đề của trang YouPass, không phải bản sao nguyên văn; vì vậy file ghi âm dùng cho đề này được tạo từ transcript (xem Audio/), KHÔNG phải audio trên trang YouPass.',i=True,sz=9,color=GREY,after=3)
    part_heading(d,'ĐÁP ÁN NHANH','20 câu · 0.5 điểm/câu')
    tb=d.add_table(rows=11,cols=4); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['Câu','Đáp án','Câu','Đáp án']): cellp(tb.rows[0].cells[j],h,b=True,sz=9.5,align='c'); shade(tb.rows[0].cells[j],LBLUE)
    for r in range(10):
        for blk in range(2):
            q=blk*10+r+1; cellp(tb.rows[r+1].cells[blk*2],f'{q}.',b=True,sz=10.5,align='c'); cellp(tb.rows[r+1].cells[blk*2+1],t['key'][q],b=True,sz=10.5,color=RED)
    widths(tb,[1.4,7.7,1.4,7.7])
    p=P(d,'',before=4,after=2); R(p,'Thang điểm: ',b=True,sz=10); R(p,'20 câu × 0,5 = 10 điểm. Part 1: viết đúng chính tả, không quá 3 từ và/hoặc số; sai chính tả hoặc vượt giới hạn từ = 0 điểm (không trừ lỗi viết hoa). Chữ số hoặc chữ đều được chấp nhận khi nghĩa đúng (xem cột "Chấp nhận").',sz=10)
    p=P(d,'',after=2); R(p,'File ghi âm: ',b=True,sz=10)
    for part in (1,2):
        add_link(p,f"../Audio/{code(n)}_Part{part}.mp3",f"{code(n)}_Part{part}.mp3",sz=10,bold=False)
        R(p,f" ({fmt_dur(DUR[str(n)][part-1]) if DUR else ''})   ",sz=10)
    part_heading(d,'PART 1.  GIẢI THÍCH CHI TIẾT',t['p1_title'])
    for p_ in t['p1']:
        q=p_['n']; qline,aline=p1_evidence(t,q)
        pp=P(d,'',before=6,after=1,keep=True); R(pp,f'Câu {q}  ',b=True,color=NAVY,sz=11); R(pp,f"[{p_['label']}]",i=True,sz=9.5,color=GREY); R(pp,'    Đáp án: ',b=True); R(pp,t['key'][q],b=True,color=RED)
        pp=P(d,'',indent=0.6,after=1,keep=True); R(pp,'Chấp nhận: ',b=True,sz=10); R(pp,accept_for(t,q),sz=10)
        pp=P(d,'',indent=0.6,after=1,keep=True); R(pp,'Bằng chứng trong băng: ',b=True,sz=10)
        if qline: R(pp,f'"{qline}"  →  ',i=True,sz=10)
        R(pp,f'"{aline}"',i=True,sz=10)
        pp=P(d,'',indent=0.6,after=2); R(pp,'Giải thích: ',b=True,sz=10); R(pp,p1_type(t,p_)+f' Câu hỏi hỏi về "{p_["label"]}" nên chú ý thông tin ngay sau câu hỏi của người dẫn thoại.',sz=10)
    part_heading(d,'PART 2.  GIẢI THÍCH CHI TIẾT – MULTIPLE CHOICE (11–15)',t['p2_title'])
    for q in t['mc']:
        a=t['key'][q['n']]
        pp=P(d,'',before=6,after=1,keep=True); R(pp,f"Câu {q['n']}  ",b=True,color=NAVY,sz=11); R(pp,'    Đáp án: ',b=True); R(pp,f"{a}. {q['opts'][a]}",b=True,color=RED)
        P(d,q['q'],i=True,sz=10,indent=0.6,after=1,keep=True)
        pp=P(d,'',indent=0.6,after=1,keep=True); R(pp,'Bằng chứng: ',b=True,sz=10); R(pp,'"'+data.sentence_of(t,q['n'])+'"',i=True,sz=10)
        pp=P(d,'',indent=0.6,after=2); R(pp,'Giải thích: ',b=True,sz=10); R(pp,explain.MC[(n,q['n'])],sz=10)
    part_heading(d,'PART 2.  GIẢI THÍCH CHI TIẾT – MATCHING (16–20)','')
    mo=dict(t['match_opts'])
    for it in t['match_items']:
        a=t['key'][it['n']]
        pp=P(d,'',before=6,after=1,keep=True); R(pp,f"Câu {it['n']}  ",b=True,color=NAVY,sz=11); R(pp,f"[{it['label']}]",i=True,sz=9.5,color=GREY); R(pp,'    Đáp án: ',b=True); R(pp,f"{a}. {mo[a]}",b=True,color=RED)
        pp=P(d,'',indent=0.6,after=1,keep=True); R(pp,'Bằng chứng: ',b=True,sz=10); R(pp,'"'+data.sentence_of(t,it['n'])+'"',i=True,sz=10)
        pp=P(d,'',indent=0.6,after=2); R(pp,'Giải thích: ',b=True,sz=10); R(pp,f'Các địa điểm/khu vực được nhắc lần lượt theo thứ tự câu hỏi 16→20; nghe từ khóa "{it["label"]}" rồi nối với vị trí đi kèm ("{mo[a]}").',sz=10)
    part_heading(d,'TRANSCRIPT (có đánh dấu vị trí đáp án)','')
    def mark(s): return re.sub(r'\[Q(\d+): ([^\]]*)\]',lambda m:f'⟦Q{m.group(1)}: {m.group(2)}⟧',s)
    P(d,'Part 1',b=True,color=NAVY,before=4,after=1)
    for l in t['script1']: P(d,mark(l),sz=9.5,after=1)
    P(d,'Part 2',b=True,color=NAVY,before=6,after=1)
    for l in t['script2']: P(d,mark(l),sz=9.5,after=2)
    d.save(path)

def quick_docx(path,T):
    from docx.enum.section import WD_ORIENT
    d=base_doc(); sec=d.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=sec.page_height,sec.page_width
    P(d,'TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN',b=True,sz=9.5,color=RED,align='c',after=2)
    P(d,'HỌC VIỆN HÀNG KHÔNG VIỆT NAM  ·  KHOA NGOẠI NGỮ',b=True,sz=10.5,color=NAVY,align='c',after=0)
    P(d,'BẢNG ĐÁP ÁN NHANH — 12 MÃ ĐỀ LISTENING',b=True,sz=17,color=NAVY,align='c',after=0)
    P(d,f'Kiểm tra giữa kỳ · {COURSE} (LISTENING) · mỗi câu 0,5đ',i=True,sz=10,color=GREY,align='c',after=6)
    tb=d.add_table(rows=13,cols=21); tb.style='Table Grid'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['Mã đề']+[str(i) for i in range(1,21)]): cellp(tb.rows[0].cells[j],h,b=True,sz=8.5,color='FFFFFF',align='c'); shade(tb.rows[0].cells[j],NAVY)
    for i,t in enumerate(T,1):
        cellp(tb.rows[i].cells[0],f"{t['n']:02d}",b=True,sz=9.5,align='c'); shade(tb.rows[i].cells[0],LBLUE)
        for j in range(1,21): cellp(tb.rows[i].cells[j],t['key'][j],sz=7.5,align='c')
    widths(tb,[1.5]+[1.1]*20)
    P(d,'Part 1 (câu 1–10): điền từ/số, tối đa 3 từ và/hoặc số. Part 2: câu 11–15 chọn A/B/C; câu 16–20 nối A–E. Xem file đáp án từng mã đề để biết cách chấp nhận đáp án và giải thích.',sz=9.5,before=4)
    d.save(path)

# ---------- HTML
CSS=G.CSS
def htop(t,teacher=False):
    n=t['n']; h=("<p class='c b red' style='font-size:9.5pt;margin:0 0 4px'>TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN</p>" if teacher else "")
    ttl=('ĐÁP ÁN VÀ GIẢI THÍCH' if teacher else 'ĐỀ KIỂM TRA GIỮA KỲ')+'  –  LISTENING'
    h+=(f"<table><tr><td width=42% class=c>BỘ XÂY DỰNG<br><b>HỌC VIỆN HÀNG KHÔNG VIỆT NAM</b><br><b>KHOA NGOẠI NGỮ</b><br>—————</td><td class=c><div class='b nav' style='font-size:16pt'>{ttl}</div>{YEAR}<br><i>Hình thức: Nghe (2 file ghi âm)</i></td></tr></table><div style='height:6px'></div>"
        f"<table class=bx><tr><td width=68%><b style='font-size:11pt'>Học phần: {COURSE}  ({COURSE_EN})</b><br>Thời gian làm bài: {TIME}</td><td class='c pk'><b style='font-size:10pt'>MÃ ĐỀ</b><br><b class=red style='font-size:22pt;line-height:1.1'>{n:02d}</b></td></tr></table>")
    return h
E=G.E
def exam_html(t):
    n=t['n']; h=f"<html><head><meta charset=utf-8>{CSS}<style>a{{color:#0563C1;font-weight:bold}}.opt{{margin-left:.7cm}}.mt td{{border:1px solid #444;text-align:center;padding:3px;background:#{LBLUE}}}</style></head><body>"+htop(t)
    h+=("<p style='margin:8px 0 2px;font-size:10pt'>Không sử dụng tài liệu [x] &nbsp;&nbsp;&nbsp; Được sử dụng tài liệu [&nbsp;&nbsp;&nbsp;] &nbsp;&nbsp;&nbsp; Nộp lại đề thi [x]</p>"
        "<p style='margin:0 0 2px;font-size:10pt'>Đề thi gồm 20 câu, nghe 2 file ghi âm (Part 1, Part 2), kèm 01 Phiếu trả lời ở trang cuối. Mỗi file chỉ được nghe MỘT LẦN. Cán bộ coi thi không giải thích gì thêm.</p><p class='b i nav' style='margin:0 0 6px;font-size:10pt'>Sinh viên ghi MÃ ĐỀ và làm bài trên PHIẾU TRẢ LỜI.</p>")
    h+=f"<table class=bx><tr><td style='background:#{LBLUE}'><b class=nav style='font-size:10.5pt'>RECORDINGS / FILE GHI ÂM</b> <i style='font-size:9.5pt'>(nhấn vào liên kết để mở file nghe)</i><br>"
    for part,qs in ((1,'Questions 1–10'),(2,'Questions 11–20')):
        h+=f"▶ <b>Part {part}</b> ({qs}): <a href='{audio_href(n,part)}'>{code(n)}_Part{part}.mp3</a> <i class=gr style='font-size:9.5pt'>(≈ {fmt_dur(DUR[str(n)][part-1]) if DUR else '2:30'} phút)</i><br>"
    h+="<i class=gr style='font-size:9pt'>Lưu ý: mở file từ thư mục \"Audio\" nằm cạnh thư mục đề thi; không đổi tên hoặc di chuyển file.</i></td></tr></table><div class=rule></div>"
    h+="<div class=ph>PART 1.  FORM COMPLETION  (Questions 1–10) <span>&nbsp; 5 points · 0.5 point each</span></div>"
    h+=f"<p class='b i nav' style='margin:0 0 4px'>{E(t['p1_directions'])} Write your answers on the ANSWER SHEET.</p>"
    h+=f"<table class=bx><tr><td colspan=2 class=c style='background:#{NAVY};color:#fff;font-weight:bold;font-size:11pt'>{E(t['p1_title'])}</td></tr>"+"".join(f"<tr style='height:26px'><td width=8% class='c b nav'>{p['n']}.</td><td>{E(p['label'])}</td></tr>" for p in t['p1'])+"</table>"
    h+="<div style='break-before:page'></div><div class=ph>PART 2.  MULTIPLE CHOICE & MATCHING  (Questions 11–20) <span>&nbsp; 5 points · 0.5 point each</span></div>"
    h+=f"<p class='b nav' style='margin:0 0 2px;font-size:11pt'>{E(t['p2_title'])}</p><p class='b i nav' style='margin:0 0 3px'>Questions 11–15.  Choose the correct letter, A, B or C.</p>"
    for q in t['mc']:
        h+=f"<div class=q><div><b class=nav>{q['n']}.</b>&nbsp; {E(q['q'])}</div>"+"".join(f"<div class=opt><b class=nav>{L}.</b>&nbsp; {E(q['opts'][L])}</div>" for L in 'ABC')+"</div>"
    h+=f"<p class='b i nav' style='margin:12px 0 4px'>{E(t['match_directions'])}  Write the correct letter, A–E, on the ANSWER SHEET.</p><table class='mt' style='width:100%'><tr>"+"".join(f"<td><b class=nav style='font-size:11pt'>{L}</b><br>{E(x)}</td>" for L,x in t['match_opts'])+"</tr></table>"
    for it in t['match_items']: h+=f"<div class=q><b class=nav>{it['n']}.</b>&nbsp; {E(it['label'])}: <span class=gr>______</span></div>"
    h+="<p class='c b nav' style='margin:14px 0 6px'>— END OF THE TEST —</p><table style='break-inside:avoid'><tr>"+"".join(f"<td class=c style='font-size:10pt'><i>TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026</i><br><b>{r}</b><br><i style='font-size:9.5pt'>(Ký và ghi rõ họ, tên)</i><div style='height:50px'></div></td>" for r in ('Trưởng Bộ môn duyệt','Giảng viên ra đề'))+"</tr></table>"
    h+="<div style='break-before:page'>"
    h+=(f"<p class='c b nav' style='margin:0'>HỌC VIỆN HÀNG KHÔNG VIỆT NAM &nbsp;·&nbsp; KHOA NGOẠI NGỮ</p><p class='c b nav' style='margin:0;font-size:17pt'>PHIẾU TRẢ LỜI &nbsp;/&nbsp; STUDENT ANSWER SHEET</p><p class='c i gr' style='margin:0 0 8px;font-size:10pt'>Kiểm tra giữa kỳ · {COURSE} (LISTENING) · {YEAR}</p>"
        f"<table class=bx><tr><td width=72%>Họ và tên: ..................................................................</td><td rowspan=3 class='c pk'><b style='font-size:10pt'>MÃ ĐỀ</b><br><b class=red style='font-size:26pt'>{n:02d}</b></td></tr><tr><td>MSSV: ................................... &nbsp;&nbsp; Lớp: ..............................</td></tr><tr><td>Phòng thi: .......... &nbsp;&nbsp; Số tờ: .......... &nbsp;&nbsp; Chữ ký SV: ....................</td></tr></table><div style='height:8px'></div>"
        f"<table class='bx c'><tr style='background:#{LBLUE}'><th>Part 1 (/5)</th><th>Part 2 (/5)</th><th>TỔNG (/10)</th><th>GV chấm 1</th><th>GV chấm 2</th></tr><tr style='height:32px'><td>......... / 5</td><td>......... / 5</td><td></td><td></td><td></td></tr></table>")
    h+="<p style='margin:12px 0 4px'><b class=nav style='font-size:11.5pt'>PART 1.</b> <i style='font-size:9.5pt'>— Write NO MORE THAN THREE WORDS AND/OR A NUMBER. Spelling must be correct.</i></p><table class='lt'><tr><th>Câu</th><th>Your answer</th><th>Câu</th><th>Your answer</th></tr>"
    for r in range(5): h+=f"<tr style='height:34px'><td width=7%><b>{r+1}.</b></td><td></td><td width=7%><b>{r+6}.</b></td><td></td></tr>"
    h+="</table><p style='margin:12px 0 4px'><b class=nav style='font-size:11.5pt'>PART 2.</b> <i style='font-size:9.5pt'>— Questions 11–15: circle ONE letter.  Questions 16–20: circle ONE letter A–E.</i></p><table class=lt><tr><th>Câu</th><th>A</th><th>B</th><th>C</th><th>Câu</th><th>A</th><th>B</th><th>C</th><th>D</th><th>E</th></tr>"
    for r in range(5):
        h+=f"<tr style='height:28px'><td><b>{11+r}.</b></td>"+"".join(f"<td class=gr>{L}</td>" for L in 'ABC')+f"<td><b>{16+r}.</b></td>"+"".join(f"<td class=gr>{L}</td>" for L in 'ABCDE')+"</tr>"
    return h+"</table></div></body></html>"
def key_html(t):
    n=t['n']; h=f"<html><head><meta charset=utf-8>{CSS}<style>a{{color:#0563C1}}</style></head><body>"+htop(t,True)
    h+=f"<p style='margin:6px 0 0'><b class=nav>Test {n}:</b> <b>{E(t['title'])}</b></p><p style='margin:0;font-size:9.5pt'><b>Nguồn tham chiếu:</b> {E(t['anchor'])} – <a href='{t['url']}'>{t['url']}</a></p>"
    h+="<p class='i gr' style='font-size:9pt;margin:2px 0 6px'>Lưu ý: file nguồn (YouPass_Listening_Parts_1_2_Teacher_Edition_12_Tests) ghi rõ câu hỏi, đáp án và transcript là bản VIẾT LẠI theo chủ đề của trang YouPass, không phải bản sao nguyên văn; vì vậy file ghi âm dùng cho đề này được tạo từ transcript (xem Audio/), KHÔNG phải audio trên trang YouPass.</p>"
    h+="<div class=ph>ĐÁP ÁN NHANH <span>&nbsp; 20 câu · 0.5 điểm/câu</span></div><table class=lt><tr><th>Câu</th><th>Đáp án</th><th>Câu</th><th>Đáp án</th></tr>"
    for r in range(10): h+=f"<tr><td width=8%><b>{r+1}.</b></td><td class='b red'>{E(t['key'][r+1])}</td><td width=8%><b>{r+11}.</b></td><td class='b red'>{E(t['key'][r+11])}</td></tr>"
    h+="</table><p style='margin:6px 0;font-size:10pt'><b>Thang điểm:</b> 20 câu × 0,5 = 10 điểm. Part 1: viết đúng chính tả, không quá 3 từ và/hoặc số; sai chính tả hoặc vượt giới hạn từ = 0 điểm (không trừ lỗi viết hoa). Chữ số hoặc chữ đều được chấp nhận khi nghĩa đúng (xem mục \"Chấp nhận\").</p>"
    h+="<p style='font-size:10pt;margin:2px 0'><b>File ghi âm:</b> "+" &nbsp; ".join(f"<a href='../Audio/{code(n)}_Part{p}.mp3'>{code(n)}_Part{p}.mp3</a> ({fmt_dur(DUR[str(n)][p-1]) if DUR else ''})" for p in (1,2))+"</p>"
    h+=f"<div class=ph>PART 1.  GIẢI THÍCH CHI TIẾT <span>&nbsp; {E(t['p1_title'])}</span></div>"
    for p_ in t['p1']:
        q=p_['n']; ql,al=p1_evidence(t,q)
        h+=(f"<div class=blk><b class=nav style='font-size:11pt'>Câu {q}</b> <i class=gr style='font-size:9.5pt'>[{E(p_['label'])}]</i> &nbsp;<b>Đáp án: <span class=red>{E(t['key'][q])}</span></b>"
            f"<div class=k style='font-size:10pt'><b>Chấp nhận:</b> {E(accept_for(t,q))}</div><div class=k style='font-size:10pt'><b>Bằng chứng trong băng:</b> <i>{('“'+E(ql)+'” → ') if ql else ''}“{E(al)}”</i></div>"
            f"<div class=k style='font-size:10pt'><b>Giải thích:</b> {E(p1_type(t,p_))} Câu hỏi hỏi về “{E(p_['label'])}” nên chú ý thông tin ngay sau câu hỏi của người dẫn thoại.</div></div>")
    h+=f"<div class=ph>PART 2.  MULTIPLE CHOICE (11–15) <span>&nbsp; {E(t['p2_title'])}</span></div>"
    for q in t['mc']:
        a=t['key'][q['n']]
        h+=(f"<div class=blk><b class=nav style='font-size:11pt'>Câu {q['n']}</b> &nbsp;<b>Đáp án: <span class=red>{a}. {E(q['opts'][a])}</span></b><div class=k style='font-size:10pt'><i>{E(q['q'])}</i></div>"
            f"<div class=k style='font-size:10pt'><b>Bằng chứng:</b> <i>“{E(data.sentence_of(t,q['n']))}”</i></div><div class=k style='font-size:10pt'><b>Giải thích:</b> {E(explain.MC[(n,q['n'])])}</div></div>")
    h+="<div class=ph>PART 2.  MATCHING (16–20)</div>"; mo=dict(t['match_opts'])
    for it in t['match_items']:
        a=t['key'][it['n']]
        h+=(f"<div class=blk><b class=nav style='font-size:11pt'>Câu {it['n']}</b> <i class=gr style='font-size:9.5pt'>[{E(it['label'])}]</i> &nbsp;<b>Đáp án: <span class=red>{a}. {E(mo[a])}</span></b>"
            f"<div class=k style='font-size:10pt'><b>Bằng chứng:</b> <i>“{E(data.sentence_of(t,it['n']))}”</i></div><div class=k style='font-size:10pt'><b>Giải thích:</b> Các địa điểm/khu vực được nhắc lần lượt theo thứ tự câu hỏi 16→20; nghe từ khóa “{E(it['label'])}” rồi nối với vị trí đi kèm (“{E(mo[a])}”).</div></div>")
    def mk(s): return re.sub(r'\[Q(\d+): ([^\]]*)\]',lambda m:f"<b class=red>⟦Q{m.group(1)}: {E(m.group(2))}⟧</b>",E(s))
    h+="<div class=ph>TRANSCRIPT <span>&nbsp; (có đánh dấu vị trí đáp án)</span></div><p class='b nav' style='margin:2px 0'>Part 1</p>"+"".join(f"<div style='font-size:9.5pt;margin:1px 0'>{mk(l)}</div>" for l in t['script1'])
    h+="<p class='b nav' style='margin:6px 0 2px'>Part 2</p>"+"".join(f"<div style='font-size:9.5pt;margin:2px 0'>{mk(l)}</div>" for l in t['script2'])
    return h+"</body></html>"
def quick_html(T):
    h=f"<html><head><meta charset=utf-8>{CSS.replace('size:Letter','size:Letter landscape')}</head><body>"
    h+=(f"<p class='c b red' style='font-size:9.5pt;margin:0'>TÀI LIỆU GIẢNG VIÊN — KHÔNG PHÁT CHO SINH VIÊN</p><p class='c b nav' style='margin:0'>HỌC VIỆN HÀNG KHÔNG VIỆT NAM &nbsp;·&nbsp; KHOA NGOẠI NGỮ</p><p class='c b nav' style='margin:0;font-size:17pt'>BẢNG ĐÁP ÁN NHANH — 12 MÃ ĐỀ LISTENING</p><p class='c i gr' style='margin:0 0 8px;font-size:10pt'>Kiểm tra giữa kỳ · {COURSE} (LISTENING) · mỗi câu 0,5đ</p>")
    h+=f"<table class=lt><tr><th style='background:#{NAVY};color:#fff'>Mã đề</th>"+"".join(f"<th style='background:#{NAVY};color:#fff'>{i}</th>" for i in range(1,21))+"</tr>"
    for t in T: h+=f"<tr><td class=b style='background:#{LBLUE}'>{t['n']:02d}</td>"+"".join(f"<td style='font-size:8pt'>{E(t['key'][j])}</td>" for j in range(1,21))+"</tr>"
    return h+"</table><p style='font-size:9.5pt'>Part 1 (câu 1–10): điền từ/số, tối đa 3 từ và/hoặc số. Part 2: câu 11–15 chọn A/B/C; câu 16–20 nối A–E. Xem file đáp án từng mã đề để biết cách chấp nhận đáp án và giải thích.</p></body></html>"

def main():
    T=data.load(); hout=os.environ.get('HTML_OUT')
    dt=os.path.join(OUT,'De_thi_va_Phieu_tra_loi'); da=os.path.join(OUT,'Dap_an'); os.makedirs(dt,exist_ok=True); os.makedirs(da,exist_ok=True)
    for t in T:
        nm=code(t['n'])
        exam_docx(os.path.join(dt,f'{nm} - DE THI.docx'),t); key_docx(os.path.join(da,f'{nm} - DAP AN.docx'),t)
        if hout:
            for sub,fn,hh in (('De_thi_va_Phieu_tra_loi',f'{nm} - DE THI',exam_html(t)),('Dap_an',f'{nm} - DAP AN',key_html(t))):
                pd=os.path.join(hout,sub); os.makedirs(pd,exist_ok=True); open(os.path.join(pd,fn+'.html'),'w',encoding='utf-8').write(hh)
    quick_docx(os.path.join(da,'00 - BANG DAP AN NHANH 12 MA DE.docx'),T)
    if hout: open(os.path.join(hout,'Dap_an','00 - BANG DAP AN NHANH 12 MA DE.html'),'w',encoding='utf-8').write(quick_html(T))

def player_html(T):
    rows="".join(f"<tr><td><b>Test {t['n']:02d}</b><br><small>{E(t['title'])}</small></td><td>Part 1 (Q1–10)<br><audio controls preload=none src='Audio/{code(t['n'])}_Part1.mp3'></audio></td><td>Part 2 (Q11–20)<br><audio controls preload=none src='Audio/{code(t['n'])}_Part2.mp3'></audio></td></tr>" for t in T)
    return ("<!doctype html><html><head><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'><title>LS1 Midterm Listening – Recordings</title>"
      "<style>body{font-family:Calibri,Arial,sans-serif;max-width:900px;margin:24px auto;padding:0 16px;color:#102A43}h1{font-size:22px}table{border-collapse:collapse;width:100%}td{border-bottom:1px solid #C9D6E5;padding:10px 6px;vertical-align:middle}small{color:#5F6368}audio{width:100%;max-width:300px}.n{background:#E8EEF4;padding:10px 14px;border-radius:8px;font-size:14px}</style></head><body>"
      "<h1>LS1 – Midterm Listening: Recordings</h1><div class=n>Mỗi file chỉ phát <b>một lần</b> khi thi. Nhấn ▶ để nghe, hoặc mở file mp3 trong thư mục <code>Audio/</code>. File ghi âm được tạo từ transcript của đề (giọng đọc tổng hợp). Mã đề trùng với tên file đề thi (LS1-L-01 … LS1-L-12).</div><br>"
      f"<table>{rows}</table></body></html>")
def finish(T):
    open(os.path.join(OUT,'00 - PLAY RECORDINGS.html'),'w',encoding='utf-8').write(player_html(T))

if __name__=="__main__":
    main(); finish(data.load())
