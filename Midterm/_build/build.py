# -*- coding: utf-8 -*-
import random,re,os,json,sys
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from prep import load,CFG
from comb import load as comb_load
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT=sys.argv[1] if len(sys.argv)>1 else 'Midterm'
SRC=sys.argv[2] if len(sys.argv)>2 else os.path.join(os.path.dirname(os.path.abspath(__file__)),'data')+'/'
BRAND="TÂM LINH"
UNI="VIETNAM AEROSPACE UNIVERSITY"
UNI_VI="Học viện Hàng không Vũ trụ Việt Nam"
NCODES=20; TIME="45 minutes"
MC_PTS=0.4; CB_PTS=0.5
BOOK="Azar & Hagen, Understanding and Using English Grammar"

# ---------------- selection ----------------
def pick_mc(tag,rng):
    d=load(tag); cfg=CFG[tag]
    pools=[[n for n in m[3] if n in d and not d[n]['excluded']] for m in cfg['modules']]
    use=[{n:0 for n in p} for p in pools]
    exams=[]
    for e in range(NCODES):
        chosen=[]
        for mi,p in enumerate(pools):
            cand=sorted(p,key=lambda n:(use[mi][n],rng.random()))[:5]
            for n in cand: use[mi][n]+=1
            chosen+=sorted(cand)
        exams.append(chosen)
    return d,exams

def comb_items(tag):
    if tag=='m1':
        items=comb_load(SRC+"Combine_Sentences__MID_GR_-L_N_1.docx.txt")
        out=[]
        for i,o in enumerate(items):
            out.append(dict(src=o['src'],ans=o['ans'],ref=f"Combine Sentences – MID GR (Lần 1), Questions 181–210, câu {181+i}"))
        return out
    A=set(tuple(o['src']) for o in comb_load(SRC+"Combine_Sentences__MID_GR_-L_N_1.docx.txt"))
    items=comb_load(SRC+"Combine_Sentences__MID_GR_-L_N_2.docx.txt")
    out=[];seen=set();secpos={}
    for o in items:
        sec=o.get('sec'); secpos[sec]=secpos.get(sec,0)+1
        k=tuple(o['src'])
        if len(k)!=3 or k in A or k in seen: continue
        if sec and sec.startswith('Questions 181'): continue
        seen.add(k)
        if sec:
            lo=int(re.match(r'Questions\s+(\d+)',sec).group(1))
            ref=f"Combine Sentences – MID GR (Lần 2), Questions {lo}–{lo+29}, câu {lo+secpos[sec]-1}"
        else:
            ref=f"Combine Sentences – MID GR (Lần 2), phần đầu file (trước nhóm 'Questions 151-180'), mục thứ {secpos[sec]}"
        out.append(dict(src=o['src'],ans=o['ans'],ref=ref))
    return out

def pick_comb(tag,rng):
    pool=comb_items(tag)
    seq=[]
    if len(pool)>=NCODES*4:
        idx=list(range(len(pool))); rng.shuffle(idx)
        return pool,[idx[i*4:(i+1)*4] for i in range(NCODES)]
    exams=[];bag=[]
    for e in range(NCODES):
        chosen=[]
        while len(chosen)<4:
            if not bag:
                bag=list(range(len(pool))); rng.shuffle(bag)
            x=bag.pop()
            if x not in chosen: chosen.append(x)
        exams.append(chosen)
    return pool,exams

def letters_plan(rng):
    for _ in range(1000):
        L=list('ABCD')*5; rng.shuffle(L)
        if all(not(L[i]==L[i+1]==L[i+2]) for i in range(18)): return L
    return L

def arrange(q,target,rng):
    correct=q['opts'][q['ans']]
    others=[v for k,v in q['opts'].items() if k!=q['ans']]
    rng.shuffle(others)
    res={};it=iter(others)
    for L in 'ABCD': res[L]=correct if L==target else next(it)
    return res

def technique(ans):
    t=[]
    if re.match(r'^[^,]{3,60}, (an? |the )?[^,]{3,}?,',ans) and not re.match(r'^[^,]+, (which|who|that)',ans): t.append('cụm đồng vị (appositive) / cụm danh từ đặt giữa hai dấu phẩy')
    if re.search(r',? which ',ans): t.append('mệnh đề quan hệ không xác định với "which"')
    if re.search(r'\bwho\b',ans): t.append('mệnh đề quan hệ với "who"')
    if re.search(r'\bthat (is|are|was|were|helps|has|have|it|they)\b',ans) or re.search(r'\w+ that [a-z]+ ',ans) and not t: t.append('mệnh đề quan hệ với "that"')
    if re.search(r'\bby \w+ing\b',ans): t.append('giới từ + V-ing (by + V-ing)')
    if re.search(r', (\w+ing|\w+ed|located|known|built|called|used|made|found) ',ans) or re.search(r'^[^,]+, \w+ing\b',ans): t.append('cụm phân từ (V-ing / V3) rút gọn mệnh đề')
    if re.search(r'\b(and|but|yet|so)\b',ans): t.append('liên từ kết hợp (and/but/yet/so) để nối các vế song song')
    if re.search(r'\b(because|since|although|though|while|when|as|if)\b',ans): t.append('mệnh đề trạng ngữ (liên từ phụ thuộc)')
    return t or ['kết hợp ý bằng cấu trúc phù hợp']

# ---------------- docx helpers ----------------
def base_doc():
    d=Document()
    s=d.styles['Normal']; s.font.name='Times New Roman'; s.font.size=Pt(11.5)
    s.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    pf=s.paragraph_format; pf.space_after=Pt(2); pf.space_before=Pt(0); pf.line_spacing=1.08
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.left_margin=sec.right_margin=Cm(2);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
    return d
def para(d,text='',bold=False,italic=False,size=None,align=None,after=None,before=None,color=None,keep=False):
    p=d.add_paragraph(); 
    if text: 
        r=p.add_run(text); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor(*color)
    if align=='c': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if align=='r': p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if after is not None: p.paragraph_format.space_after=Pt(after)
    if before is not None: p.paragraph_format.space_before=Pt(before)
    if keep: p.paragraph_format.keep_with_next=True
    return p
def run(p,text,bold=False,italic=False,size=None,color=None):
    r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    return r
def set_cell_border(cell,**kw):
    tcPr=cell._tc.get_or_add_tcPr(); b=tcPr.find(qn('w:tcBorders'))
    if b is None: b=OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right'):
        v=kw.get(edge)
        el=OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),v or 'nil'); el.set(qn('w:sz'),'6'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),'000000')
        b.append(el)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
def cell_text(cell,text,bold=False,size=None,align=None,italic=False):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if align=='c': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(1)
def footer(d,left):
    sec=d.sections[0]; p=sec.footer.paragraphs[0]; p.text=''
    r=p.add_run(left+'   |   Page '); r.font.size=Pt(9)
    for t in ('begin',None,'end'):
        if t:
            fc=OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),t); rr=p.add_run(); rr.font.size=Pt(9); rr._r.append(fc)
        else:
            it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text='PAGE'; rr=p.add_run(); rr.font.size=Pt(9); rr._r.append(it)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def header_block(d,title,midname,code,sub=None):
    t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.rows[0].cells,(Cm(8),Cm(9))): c.width=w
    l,r=t.rows[0].cells
    l.text=''; p=l.paragraphs[0]; run(p,UNI,bold=True,size=11); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p2=l.add_paragraph(); run(p2,UNI_VI,size=10.5); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r.text=''; p=r.paragraphs[0]; run(p,title,bold=True,size=13); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p2=r.add_paragraph(); run(p2,f"{midname}  –  Code: {code:02d}",bold=True,size=11); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if sub:
        p3=r.add_paragraph(); run(p3,sub,italic=True,size=10); p3.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(d,'',after=2)
def hr(d):
    p=d.add_paragraph(); pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bt=OxmlElement('w:bottom')
    bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),'6'); bt.set(qn('w:space'),'1'); bt.set(qn('w:color'),'000000'); b.append(bt); pPr.append(b)
    p.paragraph_format.space_after=Pt(4)

def add_options(d,opts):
    mx=max(len(v) for v in opts.values())
    if mx<=16:
        p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.8)); p.paragraph_format.tab_stops.add_tab_stop(Cm(8.8)); p.paragraph_format.tab_stops.add_tab_stop(Cm(12.8))
        run(p,'\t'.join(f"{L}. {opts[L]}" for L in 'ABCD'))
    elif mx<=38:
        for a_,b_ in (('A','B'),('C','D')):
            p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.8); p.paragraph_format.tab_stops.add_tab_stop(Cm(8.8))
            run(p,f"{a_}. {opts[a_]}\t{b_}. {opts[b_]}")
    else:
        for L in 'ABCD':
            p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.8); run(p,f"{L}. {opts[L]}")

def stem_par(d,num,stem):
    lines=stem.split('\n')
    p=d.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(5)
    run(p,f"{num}. ",bold=True); run(p,lines[0])
    for ln in lines[1:]:
        p2=d.add_paragraph(); p2.paragraph_format.left_indent=Cm(0.6); p2.paragraph_format.keep_with_next=True; run(p2,ln)

def answer_grid(d,answers=None):
    t=d.add_table(rows=11,cols=10); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=['No.','A','B','C','D']*2
    for j,h in enumerate(hdr):
        cell_text(t.rows[0].cells[j],h,bold=True,align='c',size=10); shade(t.rows[0].cells[j],'D9E2F3')
    for i in range(10):
        for blk in range(2):
            n=i+1+10*blk
            cell_text(t.rows[i+1].cells[blk*5],str(n),bold=True,align='c')
            for k,L in enumerate('ABCD'):
                c=t.rows[i+1].cells[blk*5+1+k]
                txt=''
                if answers: txt='●' if answers[n-1]==L else ''
                cell_text(c,txt,align='c')
        t.rows[i+1].height=Cm(0.75)
    widths=[1.1,1.3,1.3,1.3,1.3]*2
    for row in t.rows:
        for c,w in zip(row.cells,widths): c.width=Cm(w)
    return t

# ---------------- builders ----------------
def info_line(d):
    p=para(d,'',after=4)
    run(p,"Full name: ",bold=True); run(p,"…………………………………………………  ")
    run(p,"Student ID: ",bold=True); run(p,"………………………  ")
    run(p,"Class: ",bold=True); run(p,"…………………")

def build_exam(path,midno,code,qs,combs,cfgmods):
    d=base_doc(); footer(d,f"{BRAND} – Midterm {midno} – Grammar – Code {code:02d}")
    header_block(d,"MIDTERM TEST – GRAMMAR",f"Midterm {midno}",code,f"Time allowed: {TIME}  ·  24 questions")
    info_line(d)
    para(d,"Instructions: Write your answers on the STUDENT ANSWER SHEET (last page). No dictionaries or notes. Do not write on the answer sheet's reverse side.",italic=True,size=10,after=4)
    hr(d)
    para(d,f"PART I. MULTIPLE CHOICE (Questions 1–20)  –  {MC_PTS*20:.0f} points",bold=True,size=12,after=1)
    para(d,"Choose the best answer (A, B, C or D) to complete each sentence or conversation.",italic=True,after=3)
    for i,(q,arr,_) in enumerate(qs,1):
        stem_par(d,i,q['stem']); add_options(d,arr)
    hr(d)
    p=para(d,f"PART II. SENTENCE COMBINING (Questions 21–24)  –  {CB_PTS*4:.0f} points",bold=True,size=12,after=1,keep=True)
    para(d,"Combine the three sentences in each group (a, b, c) into ONE correct sentence. Do not leave out any information. You may change the word order and use relative clauses, participial phrases, appositives, conjunctions, etc.",italic=True,after=3,keep=True)
    for i,c in enumerate(combs,21):
        p=para(d,'',before=5,after=1,keep=True); run(p,f"{i}. ",bold=True)
        for j,ln in enumerate(c['src']):
            q=d.add_paragraph(); q.paragraph_format.left_indent=Cm(0.8); q.paragraph_format.keep_with_next=(j<2)
            run(q,re.sub(r'^([a-d])\.\s*',lambda m:f"{m.group(1)}.  ",ln))
    para(d,'',after=2)
    para(d,"— END OF TEST —",bold=True,align='c',before=6)
    # answer sheet
    d.add_page_break()
    header_block(d,"STUDENT ANSWER SHEET",f"Midterm {midno} – Grammar",code)
    info_line(d)
    p=para(d,'',after=4); run(p,"Date: ",bold=True); run(p,"………………………   "); run(p,"Exam code: ",bold=True); run(p,f"{code:02d}",bold=True,size=13)
    para(d,"Instructions: Use a pen. For Part I, mark ONE answer per question by putting an X (or filling the box). If you change your answer, cross out clearly and mark the new one.",italic=True,size=10,after=4)
    para(d,"PART I. MULTIPLE CHOICE (1–20)",bold=True,after=3)
    answer_grid(d)
    para(d,'',after=3)
    para(d,"PART II. SENTENCE COMBINING (21–24)",bold=True,after=3)
    for n in range(21,25):
        p=para(d,'',before=4,after=1); run(p,f"{n}. ",bold=True)
        for _ in range(2): para(d,"_"*88,after=3,before=3)
    para(d,'',after=4)
    t=d.add_table(rows=2,cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(t.rows[0].cells,["For the grader's use","Part I  (…… / 8)","Part II  (…… / 2)","TOTAL  (…… / 10)"]): cell_text(c,h,bold=True,align='c',size=10); shade(c,'EDEDED')
    for c in t.rows[1].cells: cell_text(c,'');
    t.rows[1].height=Cm(1.1)
    d.save(path)

def build_key(path,midno,code,qs,combs,cfgmods,cmods):
    d=base_doc(); footer(d,f"{BRAND} – Midterm {midno} – Grammar – Answer Key – Code {code:02d}")
    header_block(d,"ANSWER KEY & EXPLANATIONS",f"Midterm {midno} – Grammar",code,"TEACHER'S COPY – CONFIDENTIAL")
    para(d,"Đáp án nhanh – Part I (20 câu)",bold=True,size=12,after=3)
    answer_grid(d,[x[2] for x in qs])
    para(d,'',after=2)
    p=para(d,'',after=2)
    run(p,"Thang điểm: ",bold=True); run(p,f"Part I: 20 câu × {MC_PTS} = {MC_PTS*20:.0f} điểm.  Part II: 4 câu × {CB_PTS} = {CB_PTS*4:.0f} điểm (mỗi câu: 0,25 đủ cả 3 ý a–b–c; 0,25 ngữ pháp, dấu câu, không lặp thừa).  Tổng: 10 điểm.")
    p=para(d,'',after=2); run(p,"Phạm vi: ",bold=True)
    run(p,"; ".join(f"{m[0]} – {m[1]} ({m[2]})" for m in cfgmods)+". Mỗi module 5 câu trắc nghiệm (Q1–5, Q6–10, Q11–15, Q16–20).")
    hr(d)
    para(d,"PART I – GIẢI THÍCH CHI TIẾT",bold=True,size=12,after=3)
    for i,(q,arr,letter) in enumerate(qs,1):
        p=d.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(7)
        run(p,f"Câu {i}  ",bold=True); run(p,f"[{q['module']} · {q['mname']}]",italic=True,size=10); run(p,"   Đáp án: ",bold=True); run(p,f"{letter}. {arr[letter]}",bold=True,color=(0xC0,0,0))
        lines=q['stem'].split('\n')
        pq=d.add_paragraph(); pq.paragraph_format.left_indent=Cm(0.6); pq.paragraph_format.keep_with_next=True
        run(pq,' '.join(lines).replace('________','_____'),italic=True,size=10.5)
        pe=d.add_paragraph(); pe.paragraph_format.left_indent=Cm(0.6); pe.paragraph_format.keep_with_next=True
        run(pe,"Cấu trúc/quy tắc: ",bold=True,size=10.5); run(pe,q['topic']+'. ',size=10.5); run(pe,q['exp'],size=10.5)
        ps=d.add_paragraph(); ps.paragraph_format.left_indent=Cm(0.6)
        run(ps,"Nguồn: ",bold=True,size=9.5,color=(0x44,0x44,0x44))
        run(ps,f"{BOOK}, {q['chapter']} – {q['mname']}; ngân hàng câu hỏi {CFGFILE[midno]} – {q['src']}.",size=9.5,color=(0x44,0x44,0x44))
    hr(d)
    para(d,"PART II – ĐÁP ÁN THAM KHẢO & GIẢI THÍCH",bold=True,size=12,after=3,keep=True)
    para(d,"Lưu ý chấm: đây là đáp án tham khảo; chấp nhận mọi cách nối khác đúng ngữ pháp, đủ cả 3 ý và không thay đổi nghĩa.",italic=True,size=10,after=3)
    for i,c in enumerate(combs,21):
        p=d.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(7); run(p,f"Câu {i}",bold=True)
        for ln in c['src']:
            q=d.add_paragraph(); q.paragraph_format.left_indent=Cm(0.6); q.paragraph_format.keep_with_next=True; run(q,ln,size=10.5)
        pa=d.add_paragraph(); pa.paragraph_format.left_indent=Cm(0.6); pa.paragraph_format.keep_with_next=True
        run(pa,"Đáp án mẫu: ",bold=True); run(pa,c['ans'],bold=True,color=(0xC0,0,0))
        pe=d.add_paragraph(); pe.paragraph_format.left_indent=Cm(0.6); pe.paragraph_format.keep_with_next=True
        run(pe,"Giải thích: ",bold=True,size=10.5)
        run(pe,"Đáp án mẫu dùng "+"; ".join(technique(c['ans']))+". Cần giữ đủ thông tin của cả ba câu a, b, c, nối thành một câu duy nhất, tránh lặp lại chủ ngữ/đại từ (it, they, he…) không cần thiết và dùng dấu câu đúng.",size=10.5)
        ps=d.add_paragraph(); ps.paragraph_format.left_indent=Cm(0.6)
        run(ps,"Nguồn: ",bold=True,size=9.5,color=(0x44,0x44,0x44)); run(ps,c['ref']+'.',size=9.5,color=(0x44,0x44,0x44))
    d.save(path)

CFGFILE={1:'GRAMMAR_MID_L_N_1',2:'GRAMMAR_MID_L_N_2'}

def main():
    summary=[]
    for midno,tag in ((1,'m1'),(2,'m2')):
        rng=random.Random(20260000+midno)
        d,mc=pick_mc(tag,rng)
        pool,cb=pick_comb(tag,rng)
        cfgmods=CFG[tag]['modules']
        base=os.path.join(OUT,'Grammar',f'{BRAND.title() if False else BRAND} Midterm Test',f'Midterm_{midno}')
        os.makedirs(os.path.join(base,'De_thi_va_Phieu_tra_loi'),exist_ok=True)
        os.makedirs(os.path.join(base,'Dap_an'),exist_ok=True)
        for code in range(1,NCODES+1):
            plan=letters_plan(rng)
            qs=[]
            for k,n in enumerate(mc[code-1]):
                q=d[n]; arr=arrange(q,plan[k],rng); qs.append((q,arr,plan[k]))
            combs=[pool[i] for i in cb[code-1]]
            name=f"Midterm-{BRAND}_Grammar_Mid{midno}_MaDe{code:02d}"
            build_exam(os.path.join(base,'De_thi_va_Phieu_tra_loi',name+'_DeThi.docx'),midno,code,qs,combs,cfgmods)
            build_key(os.path.join(base,'Dap_an',name+'_DapAn.docx'),midno,code,qs,combs,cfgmods,None)
            summary.append(dict(mid=midno,code=code,mc=[q['id'] for q,_,_ in qs],ans=''.join(a for _,_,a in qs),comb=cb[code-1]))
    json.dump(summary,open(os.path.join(OUT,'_build','exam_composition.json'),'w'),ensure_ascii=False,indent=1)
if __name__=="__main__": main()
